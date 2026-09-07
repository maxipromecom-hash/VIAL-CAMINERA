import os
import re
import sys
import json
import time
import shutil
import threading
import unicodedata
from datetime import datetime, date
from difflib import SequenceMatcher
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter.scrolledtext import ScrolledText

from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.common.exceptions import TimeoutException, WebDriverException

VERSION = "20.0"
URL_LOGIN = "https://sistemas.seguridad.mendoza.gov.ar/vialcaminera//servlet/com.ktksuitelr.mdlsgt.hlogin2"
URL_CONSULTA_HINT = "wpconsultaantecedentes"

# v13 UNIVERSAL: la estructura del Excel se detecta automáticamente.
# Ya no dependemos de posiciones fijas de columnas ni de una hoja específica.
# Por decisión del usuario NO se comparan Fecha, Motor ni Chasis.
CAMPOS_COMPARADOS = ("DOMINIO↔ACTA", "MARCA/MODELO", "COLOR")



# Diccionario editable de normalización Marca/Modelo.
DICCIONARIO_DEFAULT = {
    "MOTO MEL": "MOTOMEL",
    "MOTOMEL": "MOTOMEL",
    "BUSSINES": "BUSINESS",
    "BUSINES": "BUSINESS",
    "BUSSINESS": "BUSINESS",
    "ZANELLA SOL BUSSINES": "ZANELLA SOL BUSINESS",
    "ZANELLA SOL BUSINES": "ZANELLA SOL BUSINESS",
    "HONDA WAVE 110": "HONDA WAVE",
    "HONDA WAVE110": "HONDA WAVE",
    "CORVEN ENERGY 110": "CORVEN ENERGY",
    "GILERA SMASH 110": "GILERA SMASH",
    "YAMAHA YBR 125": "YAMAHA YBR"
}

def directorio_app():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def cargar_diccionario_modelos():
    ruta = os.path.join(directorio_app(), "diccionario_vehiculos.json")
    data = dict(DICCIONARIO_DEFAULT)
    try:
        if os.path.exists(ruta):
            with open(ruta, "r", encoding="utf-8") as f:
                extra = json.load(f)
            if isinstance(extra, dict):
                data.update({norm(k): norm(v) for k, v in extra.items()})
        else:
            with open(ruta, "w", encoding="utf-8") as f:
                json.dump(DICCIONARIO_DEFAULT, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    return {norm(k): norm(v) for k, v in data.items()}

DICCIONARIO_MODELOS = None

def aplicar_diccionario_modelo(v):
    global DICCIONARIO_MODELOS
    if DICCIONARIO_MODELOS is None:
        DICCIONARIO_MODELOS = cargar_diccionario_modelos()
    n = normalizar_modelo(v)
    # Sustituciones de frase completa y luego parciales.
    if n in DICCIONARIO_MODELOS:
        n = DICCIONARIO_MODELOS[n]
    for k, val in sorted(DICCIONARIO_MODELOS.items(), key=lambda kv: len(kv[0]), reverse=True):
        if k and k in n:
            n = n.replace(k, val)
    return norm(n)

def combinar_marca_modelo(datos):
    marca = str(datos.get("Marca", "") or "").strip()
    modelo = str(datos.get("Modelo", "") or "").strip()
    # Muchos registros GeneXus usan 0 como modelo vacío.
    if norm(modelo) in {"", "0", "S D", "SD", "SIN MODELO"}:
        modelo = ""
    nm = norm(marca)
    nmod = norm(modelo)
    if nmod and nmod not in nm:
        return (marca + " " + modelo).strip()
    return marca or modelo

def norm(v):
    if v is None:
        return ""
    s = str(v).upper().strip()
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def compact(v):
    return re.sub(r"[^A-Z0-9]", "", norm(v))


def xpath_literal(s):
    s = str(s)
    if "'" not in s:
        return f"'{s}'"
    if '"' not in s:
        return f'"{s}"'
    parts = s.split("'")
    return "concat(" + ", \"'\", ".join(f"'{p}'" for p in parts) + ")"


def visible(el):
    try:
        return el.is_displayed()
    except Exception:
        return False


def fecha_excel(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    s = norm(v)
    m = re.search(r"(\d{1,2})[ /.-](\d{1,2})[ /.-](\d{2,4})", s)
    if m:
        d, mo, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        return f"{int(d):02d}/{int(mo):02d}/{y}"
    return str(v or "").strip()


def normalizar_tipo(v):
    n = norm(v)
    aliases = {
        "MOTO": "MOTOCICLETA",
        "MOTOCICLETA": "MOTOCICLETA",
        "MOTOVEHICULO": "MOTOCICLETA",
        "AUTO": "AUTOMOVIL",
        "AUTOMOVIL": "AUTOMOVIL",
        "AUTOMOTOR": "AUTOMOVIL",
        "CAMIONETA": "CAMIONETA",
        "CAMION": "CAMION",
        "BICIMOTO": "BICIMOTO",
    }
    return aliases.get(n, n)


def normalizar_color(v):
    n = norm(v)
    cambios = {
        "NEGRA": "NEGRO", "NEGRO": "NEGRO",
        "BLANCA": "BLANCO", "BLANCO": "BLANCO",
        "ROJA": "ROJO", "ROJO": "ROJO",
        "VERDE": "VERDE", "GRIS": "GRIS",
        "AZUL": "AZUL", "AMARILLA": "AMARILLO", "AMARILLO": "AMARILLO",
    }
    toks = [cambios.get(t, t) for t in n.split()]
    return " ".join(toks)


def dominio_sin_chapa(v):
    n = norm(v)
    return n in {
        "", "S D", "SD", "S CHAPA", "SIN CHAPA", "SIN DOMINIO",
        "NO POSEE", "NO POSEE DOMINIO", "NO SE DIVISA", "NO SE OBSERVA"
    }



def extraer_dominio_real(v):
    """Devuelve una patente utilizable para búsqueda o None si realmente no hay dominio.

    Conserva la letra inicial de dominios Mercosur de motos (ej. A073JXE) y
    soporta formatos con guiones/espacios y observaciones como
    'S/CHAPA (A032JLX)' o 'S/DOMINIO (537KEO)'.
    """
    raw = str(v or "").upper()

    # Buscar sobre el texto original, permitiendo sólo espacios/guiones DENTRO
    # del dominio. No compactamos toda la frase antes de buscar porque eso puede
    # unir letras de palabras como "S/CHAPA" con la patente entre paréntesis.
    patrones = [
        r"(?<![A-Z0-9])[A-Z]{2}[\s-]*\d{3}[\s-]*[A-Z]{2}(?![A-Z0-9])",  # AA123AA
        r"(?<![A-Z0-9])[A-Z][\s-]*\d{3}[\s-]*[A-Z]{3}(?![A-Z0-9])",   # A123ABC
        r"(?<![A-Z0-9])[A-Z]{3}[\s-]*\d{3}(?![A-Z0-9])",               # ABC123
        r"(?<![A-Z0-9])\d{3}[\s-]*[A-Z]{3}(?![A-Z0-9])",               # 123ABC
    ]

    for patron in patrones:
        m = re.search(patron, raw)
        if m:
            return re.sub(r"[^A-Z0-9]", "", m.group(0))

    if dominio_sin_chapa(v) or any(x in norm(v) for x in ("S DOMINIO", "S CHAPA", "SIN DOMINIO", "SIN CHAPA")):
        return None
    return None


def candidatos_dominio_verificacion(v):
    """Genera variantes de dominio para la consulta.

    Para dominios clásicos de seis caracteres prueba ambas orientaciones:
    948HPD -> 948HPD y HPD948; HPD948 -> HPD948 y 948HPD.
    Los dominios Mercosur de siete caracteres (A123ABC / AA123AA) se conservan
    tal cual porque invertirlos produciría un formato inválido.
    """
    d = extraer_dominio_real(v)
    if not d:
        return []
    out=[d]
    if re.fullmatch(r"\d{3}[A-Z]{3}", d):
        out.append(d[3:] + d[:3])
    elif re.fullmatch(r"[A-Z]{3}\d{3}", d):
        out.append(d[3:] + d[:3])
    # Quitar duplicados preservando orden.
    return list(dict.fromkeys(out))


# ---------------- v17: DETECCIÓN UNIVERSAL DE PLANILLAS ----------------
COLORES_BASE = {
    "NEGRO", "BLANCO", "ROJO", "GRIS", "AZUL", "VERDE", "AMARILLO",
    "BORDO", "BORDEAUX", "NARANJA", "CELESTE", "MARRON", "BEIGE",
    "VIOLETA", "DORADO", "PLATEADO", "CREMA"
}
TIPOS_BASE = {
    "MOTO", "MOTOCICLETA", "MOTOVEHICULO", "AUTO", "AUTOMOVIL", "AUTOMOTOR",
    "CAMIONETA", "CAMION", "BICIMOTO", "CUATRICICLO", "TRAILER", "ACOPLADO"
}


def _header(v):
    return norm(v)


def _parece_color(v):
    n = normalizar_color(v)
    if not n:
        return False
    toks = set(n.split())
    return bool(toks & COLORES_BASE)


def _parece_tipo(v):
    return norm(v) in TIPOS_BASE or normalizar_tipo(v) in {"MOTOCICLETA", "AUTOMOVIL", "CAMIONETA", "CAMION", "BICIMOTO"}


def _parece_marca_modelo(v):
    n = norm(v)
    if not n or len(n) < 3:
        return False
    if _parece_color(v) or _parece_tipo(v) or extraer_dominio_real(v):
        return False
    # Texto típico de marca/modelo: letras y opcionalmente cilindrada/modelo.
    return bool(re.search(r"[A-Z]{3,}", n))




def acta_canonica(valor):
    """Normaliza actas equivalentes ignorando ceros a la izquierda tras el prefijo.

    Ej.: L0002302941 == L2302941; X000123456 == X123456.
    """
    if valor is None:
        return ""
    c = compact(valor)
    m = re.fullmatch(r"([A-Z])(\d{6,12})", c)
    if not m:
        return c
    pref, dig = m.groups()
    dig = dig.lstrip("0") or "0"
    return pref + dig


def cuerpo_contiene_acta_equivalente(texto, acta):
    objetivo = acta_canonica(acta)
    if not objetivo:
        return False
    cuerpo = compact(texto)
    if objetivo in cuerpo:
        return True
    # Extrae identificadores L/X del resultado y compara en forma canónica.
    for pref, dig in re.findall(r"([LXRF])\s*[- ]?\s*(\d{6,12})", str(texto or "").upper()):
        if acta_canonica(pref + dig) == objetivo:
            return True
    return False

def _digitos_acta_sin_prefijo(valor):
    """Devuelve los dígitos cuando la fila trae un acta vial sin letra.

    Acepta, por ejemplo: 7504190, AV 7504190, A.V. 7504190, ACTA 7504190.
    No confunde sumarios cortos del tipo 16/19 porque exige entre 6 y 12 dígitos.
    """
    if valor is None:
        return None
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            if float(valor).is_integer():
                dig = str(int(valor))
                return dig if 6 <= len(dig) <= 12 else None
        except Exception:
            return None
    raw = str(valor).upper().strip()
    if not raw:
        return None
    # Si ya tiene prefijo explícito L/X/R/F, no es un acta sin letra.
    if re.search(r"\b[LXRF]\s*[- ]?\s*\d{6,12}\b", raw):
        return None
    # Celda compuesta sólo por el número.
    m = re.fullmatch(r"\s*(\d{6,12})\s*", raw)
    if m:
        return m.group(1)
    # Formatos frecuentes: AV 7504190 / A.V. 7504190 / ACTA 7504190 / ACTA VIAL 7504190.
    m = re.search(r"\b(?:A\.?\s*V\.?|ACTA(?:\s+VIAL)?|VIAL)\s*[:\-]?\s*(\d{6,12})\b", raw)
    if m:
        return m.group(1)
    return None


def candidatos_acta_verificacion(acta, acta_original=None):
    """Genera los identificadores que deben probarse en el sistema.

    Si la fuente no trae letra, prueba en este orden: L, X, R y F.
    Si ya existe una letra explícita, conserva únicamente esa acta.
    """
    dig = _digitos_acta_sin_prefijo(acta_original)
    if dig:
        return [f"L{dig}", f"X{dig}", f"R{dig}", f"F{dig}"]
    a = compact(acta)
    if re.fullmatch(r"[LXRF]\d{6,12}", a):
        return [a]
    return [acta] if acta else []


def extraer_acta_vial(valor):
    """Extrae el identificador vial útil desde formatos heterogéneos.

    Ejemplos:
      ACTA L2401794 -> L2401794
      L0002301916 -> L0002301916
      SUM 163/19 ACTA X5300401 -> X5300401
      4300411 (San Cristóbal) -> L4300411
    Un SUMARIO/EXPTE sin identificador vial explícito devuelve None.
    """
    if valor is None:
        return None
    # Números puros de la columna EXTE-ACTA VIAL de San Cristóbal.
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        try:
            if float(valor).is_integer():
                dig = str(int(valor))
                if 6 <= len(dig) <= 12:
                    return "L" + dig
        except Exception:
            pass
    raw = str(valor).upper().strip()
    if not raw or norm(raw) in {"NO CONSTA", "SIN DATOS", "S D", "SD", "NINGUNO"}:
        return None
    # Identificador vial explícito L/X/R/F + números, aun si está dentro de una frase.
    m = re.search(r"\b([LXRF])\s*[- ]?\s*(\d{6,12})\b", raw)
    if m:
        return m.group(1) + m.group(2)
    # Cualquier letra + 6 o más dígitos si viene precedido por ACTA.
    m = re.search(r"\bACTA\s*[:\-]?\s*([A-Z])\s*[- ]?\s*(\d{6,12})\b", raw)
    if m:
        return m.group(1) + m.group(2)
    # Cuando el acta no trae letra, usamos L sólo como representación inicial.
    # La verificación real probará L, X, R y F mediante candidatos_acta_verificacion().
    dig = _digitos_acta_sin_prefijo(valor)
    if dig:
        return "L" + dig
    return None


def _score_header_row(ws, fila):
    score = 0
    encontrados = set()
    for c in range(1, ws.max_column + 1):
        h = _header(ws.cell(fila, c).value)
        if not h:
            continue
        if h == "TIPO" or "TIPO VEHIC" in h: encontrados.add("tipo")
        if "MARCA" in h and "MODELO" in h: encontrados.add("marca")
        if "COLOR" in h: encontrados.add("color")
        if "DOMINIO" in h or "PATENTE" in h: encontrados.add("dominio")
        if "ACTA" in h or "SUMARIO" in h or ("EXTE" in h and "VIAL" in h): encontrados.add("acta")
        if h in {"OBS", "OBSERVACIONES", "VERIFICACION VIAL"}: encontrados.add("obs")
        if "INTERVIENE" in h or "INTERVINIENTE" in h: encontrados.add("interviene")
        if "FECHA" in h: encontrados.add("fecha")
        if "DEPENDENCIA" in h: encontrados.add("dependencia")
        if "INGRESO" in h or "REG INTERNO" in h: encontrados.add("registro")
    # Los cinco campos básicos pesan más; extras ayudan a elegir entre hojas duplicadas.
    score = sum(3 if x in encontrados else 0 for x in ("tipo", "marca", "color", "dominio", "acta"))
    score += sum(1 for x in encontrados if x not in {"tipo", "marca", "color", "dominio", "acta"})
    return score, encontrados


def _sample_values(ws, header_row, col, max_n=90):
    vals=[]
    fin=min(ws.max_row, header_row + max_n)
    for r in range(header_row + 1, fin + 1):
        v=ws.cell(r,col).value
        if v is not None and str(v).strip():
            vals.append(v)
    return vals


def _ratio(vals, pred):
    if not vals:
        return 0.0
    ok=0
    for v in vals:
        try:
            if pred(v): ok += 1
        except Exception:
            pass
    return ok / max(1, len(vals))


def _choose_semantic_col(ws, header_row, semantic):
    best=None
    for c in range(1, ws.max_column + 1):
        h=_header(ws.cell(header_row,c).value)
        vals=_sample_values(ws,header_row,c)
        hb=0.0
        data=0.0
        if semantic == "dominio":
            hb = 2.5 if ("DOMINIO" in h or "PATENTE" in h) else 0.0
            data = _ratio(vals, lambda v: bool(extraer_dominio_real(v)) or dominio_sin_chapa(v) or "S DOMINIO" in norm(v)) * 10
        elif semantic == "color":
            hb = 2.5 if "COLOR" in h else 0.0
            data = _ratio(vals, _parece_color) * 10
        elif semantic == "tipo":
            hb = 4.0 if (h == "TIPO" or "TIPO VEHIC" in h) else 0.0
            data = _ratio(vals, _parece_tipo) * 8
        elif semantic == "marca_modelo":
            hb = 8.0 if ("MARCA" in h and "MODELO" in h) else 0.0
            data = _ratio(vals, _parece_marca_modelo) * 2
        elif semantic == "acta":
            hb = 8.0 if ("ACTA" in h or "SUMARIO" in h or ("EXTE" in h and "VIAL" in h)) else 0.0
            data = _ratio(vals, lambda v: extraer_acta_vial(v) is not None) * 4
        elif semantic == "interviene":
            hb = 6.0 if ("INTERVIENE" in h or "INTERVINIENTE" in h) else 0.0
            data = 0.0
        score=hb+data
        if best is None or score > best[0]:
            best=(score,c,h)
    return best[1] if best and best[0] >= 2.0 else None


def detectar_estructura_excel(wb):
    """Selecciona automáticamente hoja, fila de encabezados y columnas útiles."""
    candidatos=[]
    for ws in wb.worksheets:
        max_head=min(12, ws.max_row)
        for r in range(1,max_head+1):
            score, encontrados=_score_header_row(ws,r)
            if score >= 10:
                # Una hoja más completa y con más filas gana en caso de empate.
                candidatos.append((score, len(encontrados), min(ws.max_row,1000), ws, r))
    if not candidatos:
        raise RuntimeError("No pude detectar una tabla con TIPO, MARCA/MODELO, COLOR, DOMINIO y ACTA/SUMARIO.")
    candidatos.sort(key=lambda x:(x[0],x[1],x[2]), reverse=True)
    _,_,_,ws,header_row=candidatos[0]

    cfg={
        "sheet": ws,
        "sheet_name": ws.title,
        "header_row": header_row,
        "tipo": _choose_semantic_col(ws,header_row,"tipo"),
        "marca_modelo": _choose_semantic_col(ws,header_row,"marca_modelo"),
        "color": _choose_semantic_col(ws,header_row,"color"),
        "dominio": _choose_semantic_col(ws,header_row,"dominio"),
        "acta": _choose_semantic_col(ws,header_row,"acta"),
        "interviene": _choose_semantic_col(ws,header_row,"interviene"),
    }
    if not all(cfg.get(k) for k in ("marca_modelo","color","dominio","acta")):
        raise RuntimeError("Detecté la hoja, pero faltan columnas esenciales (Marca/Modelo, Color, Dominio o Acta).")

    # Usar una columna dedicada si ya existe. OBS simple puede contener datos operativos,
    # por eso NO se pisa; se crea VERIFICACION VIAL al final.
    col_ver=None
    for c in range(1,ws.max_column+1):
        h=_header(ws.cell(header_row,c).value)
        if h in {"VERIFICACION VIAL", "OBSERVACIONES"}:
            col_ver=c
            break
    if col_ver is None:
        col_ver=ws.max_column+1
        ws.cell(header_row,col_ver).value="VERIFICACION VIAL"
        try:
            from copy import copy
            src=ws.cell(header_row,max(1,col_ver-1))
            dst=ws.cell(header_row,col_ver)
            if src.has_style:
                dst._style=copy(src._style)
                dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)
                dst.alignment=copy(src.alignment); dst.number_format=src.number_format
        except Exception:
            pass
        try:
            ws.column_dimensions[ws.cell(header_row,col_ver).column_letter].width=30
        except Exception:
            pass
    cfg["verificacion"]=col_ver
    return cfg


def extraer_acta_de_fila(ws, fila, cfg):
    """Usa la columna detectada y, si hace falta, busca un L/X/R/F######## en toda la fila."""
    raw=ws.cell(fila,cfg["acta"]).value if cfg.get("acta") else None
    acta=extraer_acta_vial(raw)
    if acta:
        return acta, raw
    # Valle de Uco puede tener NO CONSTA en SUMARIO pero un QRU L... en OBS.
    for c in range(1,ws.max_column+1):
        v=ws.cell(fila,c).value
        if v is None: continue
        m=re.search(r"\b([LXRF])\s*[- ]?\s*(\d{6,12})\b", str(v).upper())
        if m:
            return m.group(1)+m.group(2), raw
    return None, raw


def es_acta_vial(acta):
    """Identificador vial explícito utilizable en la consulta."""
    return bool(re.fullmatch(r"[LXRF]\d{6,12}", compact(acta)))


def extraer_info_objeto_simple(driver):
    """Lee directamente la tabla de Información de Objeto, sin entrar a la lupa.

    La ventana mostrada por el sistema ya contiene Modelo, Marca y Color. Esta
    función mapea encabezados -> primera fila visible y evita el modal de Datos
    Descriptivos que estaba devolviendo AUTOMOVIL/vacíos de forma repetida.
    """
    if not _activar_contexto_info_objeto(driver):
        return {}
    try:
        tablas = driver.find_elements(By.XPATH, "//table")
    except Exception:
        tablas=[]
    for tabla in tablas:
        try:
            if not visible(tabla):
                continue
            filas=[r for r in tabla.find_elements(By.XPATH, ".//tr") if visible(r)]
            if len(filas) < 2:
                continue
            # buscar una fila de encabezados que contenga al menos Marca y Color
            for idx, hr in enumerate(filas[:-1]):
                hs=[norm(c.text) for c in hr.find_elements(By.XPATH, "./th|./td")]
                if "MARCA" not in hs or "COLOR" not in hs:
                    continue
                # primera fila posterior con contenido
                for dr in filas[idx+1:]:
                    celdas=dr.find_elements(By.XPATH, "./td|./th")
                    vals=[(c.text or "").strip() for c in celdas]
                    if not any(vals):
                        continue
                    datos={}
                    for i,h in enumerate(hs):
                        if i >= len(vals):
                            break
                        if h == "MARCA": datos["Marca"] = vals[i]
                        elif h == "MODELO": datos["Modelo"] = vals[i]
                        elif h == "COLOR": datos["Color"] = vals[i]
                        elif h == "DOMINIO": datos["Dominio"] = vals[i]
                    if datos.get("Marca") or datos.get("Modelo") or datos.get("Color"):
                        return datos
        except Exception:
            continue
    return {}

def normalizar_modelo(v):
    n = norm(v)
    reemplazos = {
        "BUSSINES": "BUSINESS",
        "BUSINES": "BUSINESS",
        "BUSSINESS": "BUSINESS",
        "MOTOMEL": "MOTOMEL",
    }
    toks = []
    for t in n.split():
        t = reemplazos.get(t, t)
        # La cilindrada no debe provocar una diferencia si el sistema no la muestra.
        if re.fullmatch(r"\d{2,4}CC", t) or t == "CC":
            continue
        toks.append(t)
    return " ".join(toks)


def tokens_significativos(v):
    stop = {"DE", "DEL", "LA", "EL", "LOS", "LAS", "Y", "N", "NRO", "NUMERO", "CALLE", "AV", "AVENIDA"}
    return [t for t in norm(v).split() if len(t) >= 2 and t not in stop]


def coincide_fecha(excel, web):
    f = fecha_excel(excel)
    return bool(f) and f in str(web or "")


def coincide_tipo(excel, web):
    return normalizar_tipo(excel) == normalizar_tipo(web) or compact(normalizar_tipo(excel)) in compact(normalizar_tipo(web))


def coincide_color(excel, web):
    a = normalizar_color(excel)
    b = normalizar_color(web)
    if not a:
        return True
    ta = set(a.split())
    tb = set(b.split())
    return bool(ta) and ta.issubset(tb)


def coincide_dominio(excel, web):
    if dominio_sin_chapa(excel):
        return dominio_sin_chapa(web)
    a = compact(excel)
    b = compact(web)
    return bool(a) and (a == b or a in b or b in a)


def coincide_marca_modelo(excel, web):
    a = aplicar_diccionario_modelo(excel)
    b = aplicar_diccionario_modelo(web)
    if not a:
        return True
    if not b:
        return False
    ca, cb = compact(a), compact(b)
    if ca == cb or ca in cb or cb in ca:
        return True
    ta = tokens_significativos(a)
    tb = tokens_significativos(b)
    if not ta:
        return True
    # Coincidencia por variables/palabras significativas.
    similares = 0
    for t in ta:
        if t in tb or any(SequenceMatcher(None, t, w).ratio() >= 0.78 for w in tb):
            similares += 1
    # La marca suele ser la primera variable y debe coincidir.
    marca_ok = (ta[0] in tb) or any(SequenceMatcher(None, ta[0], w).ratio() >= 0.86 for w in tb)
    return marca_ok and similares / len(ta) >= 0.60

def acronimo(v):
    toks = tokens_significativos(v)
    return "".join(t[0] for t in toks if t)


def coincide_interviene(excel, web):
    a = norm(excel)
    b = norm(web)
    if not a:
        return True
    ca, cb = compact(a), compact(b)
    if ca and ca in cb:
        return True

    # U.R.V.G.A. equivale a U.R.V. GENERAL ALVEAR.
    # Se compara también por iniciales de las palabras de la dependencia.
    ac_web = acronimo(b)
    ac_excel = acronimo(a)
    if ca and (ca in ac_web or ac_web in ca):
        return True
    if ac_excel and (ac_excel in ac_web or ac_web in ac_excel):
        return True

    ta = tokens_significativos(a)
    tb = tokens_significativos(b)
    if not ta:
        return True
    presentes = sum(1 for t in ta if t in tb)
    return presentes / len(ta) >= 0.60


def find_input_acta(driver, log=None):
    """Localiza el input REAL de Nro Acta usando proximidad visual al rótulo.

    En este sistema GeneXus hay contenedores cuyo texto completo también contiene
    "Nro Acta". La v3 podía tomar un input equivocado. Esta versión busca primero
    elementos hoja con el rótulo y elige el campo de texto visible más cercano,
    a la derecha y en la misma fila.
    """
    def _log(msg):
        if log:
            try:
                log(msg)
            except Exception:
                pass

    # 1) Reunir inputs de texto visibles y habilitados.
    inputs = []
    for e in driver.find_elements(By.CSS_SELECTOR, "input"):
        try:
            tipo = (e.get_attribute("type") or "text").lower()
            if tipo not in {"text", "search", "tel", "number", ""}:
                continue
            if not visible(e) or not e.is_enabled():
                continue
            r = e.rect
            if r.get("width", 0) <= 10 or r.get("height", 0) <= 5:
                continue
            inputs.append(e)
        except Exception:
            pass

    # 2) Buscar rótulos hoja cuyo texto sea exactamente Nro Acta (con o sin ':').
    labels = []
    xpath_labels = (
        "//*[not(*) and string-length(normalize-space(.)) > 0 "
        "and contains(translate(normalize-space(.),"
        "'abcdefghijklmnopqrstuvwxyzáéíóúñ','ABCDEFGHIJKLMNOPQRSTUVWXYZÁÉÍÓÚÑ'),'NRO ACTA')]"
    )
    for et in driver.find_elements(By.XPATH, xpath_labels):
        try:
            txt = norm(et.text)
            if txt in {"NRO ACTA", "NRO ACTA N", "N ACTA", "NUMERO ACTA"} or txt.startswith("NRO ACTA"):
                if visible(et):
                    labels.append(et)
        except Exception:
            pass

    # 3) Seleccionar por geometría: misma altura y a la derecha del rótulo.
    mejores = []
    for lab in labels:
        try:
            lr = lab.rect
            ly = lr.get("y", 0) + lr.get("height", 0) / 2
            lx_right = lr.get("x", 0) + lr.get("width", 0)
            for inp in inputs:
                ir = inp.rect
                iy = ir.get("y", 0) + ir.get("height", 0) / 2
                ix = ir.get("x", 0)
                dy = abs(iy - ly)
                dx = ix - lx_right
                # Debe estar aproximadamente en la misma fila y no muy lejos a la izquierda.
                if dy <= 32 and dx >= -20:
                    score = dy * 10 + max(dx, 0)
                    mejores.append((score, inp, lab))
        except Exception:
            pass

    if mejores:
        mejores.sort(key=lambda t: t[0])
        inp = mejores[0][1]
        try:
            _log("  Campo detectado: id='{}' name='{}' x={} y={}".format(
                inp.get_attribute("id") or "",
                inp.get_attribute("name") or "",
                round(inp.rect.get("x", 0), 1),
                round(inp.rect.get("y", 0), 1),
            ))
        except Exception:
            pass
        return inp

    # 4) Fallback por atributos GeneXus / nombre del control.
    candidatos = []
    for e in inputs:
        try:
            attrs = " ".join([
                e.get_attribute("id") or "",
                e.get_attribute("name") or "",
                e.get_attribute("placeholder") or "",
                e.get_attribute("title") or "",
            ])
            na = norm(attrs)
            if "ACTA" in na and "DOCUMENT" not in na:
                candidatos.append(e)
        except Exception:
            pass
    if candidatos:
        return candidatos[0]

    # 5) Último fallback: en la pantalla conocida, Nro Acta es el input de texto
    # de la columna izquierda que está debajo de Nro Documento. Elegimos por Y.
    if len(inputs) >= 2:
        try:
            # Agrupar por mitad izquierda de la pantalla y ordenar verticalmente.
            ancho = driver.execute_script("return window.innerWidth || document.documentElement.clientWidth") or 1400
            izquierdos = [e for e in inputs if e.rect.get("x", 0) < ancho * 0.55]
            izquierdos.sort(key=lambda e: e.rect.get("y", 0))
            if len(izquierdos) >= 2:
                return izquierdos[-1]
        except Exception:
            pass
    return None

def find_input_dominio(driver, log=None):
    def _log(msg):
        if log:
            try: log(msg)
            except Exception: pass
    inputs=[]
    for e in driver.find_elements(By.CSS_SELECTOR, "input"):
        try:
            tipo=(e.get_attribute("type") or "text").lower()
            if tipo not in {"text","search","tel","number",""} or not visible(e) or not e.is_enabled():
                continue
            if e.rect.get("width",0)>10 and e.rect.get("height",0)>5:
                inputs.append(e)
        except Exception: pass
    labels=[]
    xp="//*[not(*) and contains(translate(normalize-space(.),'abcdefghijklmnopqrstuvwxyzáéíóúñ','ABCDEFGHIJKLMNOPQRSTUVWXYZÁÉÍÓÚÑ'),'DOMINIO')]"
    for lab in driver.find_elements(By.XPATH,xp):
        try:
            if visible(lab) and norm(lab.text) in {"DOMINIO","NRO DOMINIO","N DOMINIO"}: labels.append(lab)
        except Exception: pass
    cand=[]
    for lab in labels:
        lr=lab.rect; ly=lr.get('y',0)+lr.get('height',0)/2; rx=lr.get('x',0)+lr.get('width',0)
        for inp in inputs:
            ir=inp.rect; iy=ir.get('y',0)+ir.get('height',0)/2; dx=ir.get('x',0)-rx; dy=abs(iy-ly)
            if dy<=35 and dx>=-20: cand.append((dy*10+max(dx,0),inp))
    if cand:
        cand.sort(key=lambda x:x[0]); e=cand[0][1]
        _log("  Campo Dominio detectado: id='{}' name='{}'".format(e.get_attribute('id') or '',e.get_attribute('name') or ''))
        return e
    for e in inputs:
        attrs=norm(' '.join([e.get_attribute('id') or '',e.get_attribute('name') or '',e.get_attribute('placeholder') or '']))
        if 'DOMINIO' in attrs: return e
    return None

def cargar_input_geneXus(driver, inp, valor):
    try:
        driver.execute_script("arguments[0].scrollIntoView({block:'center'});", inp)
        inp.click(); inp.send_keys(Keys.CONTROL,'a'); inp.send_keys(Keys.BACKSPACE)
        for ch in str(valor):
            inp.send_keys(ch); time.sleep(0.035)
        inp.send_keys(Keys.TAB); time.sleep(0.25)
        actual=(inp.get_attribute('value') or '').strip()
        if compact(actual)==compact(valor): return True
        driver.execute_script("""const el=arguments[0],val=arguments[1]; const p=Object.getPrototypeOf(el); const d=Object.getOwnPropertyDescriptor(p,'value'); if(d&&d.set)d.set.call(el,val); else el.value=val; ['input','change','keyup','blur'].forEach(t=>el.dispatchEvent(new Event(t,{bubbles:true})));""", inp, str(valor))
        time.sleep(0.25)
        return compact(inp.get_attribute('value') or '')==compact(valor)
    except Exception:
        return False

def click_buscar(driver):
    xpaths = [
        "//button[contains(translate(normalize-space(.),'abcdefghijklmnopqrstuvwxyz','ABCDEFGHIJKLMNOPQRSTUVWXYZ'),'BUSCAR')]",
        "//input[(@type='button' or @type='submit') and contains(translate(@value,'abcdefghijklmnopqrstuvwxyz','ABCDEFGHIJKLMNOPQRSTUVWXYZ'),'BUSCAR')]",
        "//a[contains(translate(normalize-space(.),'abcdefghijklmnopqrstuvwxyz','ABCDEFGHIJKLMNOPQRSTUVWXYZ'),'BUSCAR')]",
    ]
    for xp in xpaths:
        for e in driver.find_elements(By.XPATH, xp):
            if visible(e) and e.is_enabled():
                driver.execute_script("arguments[0].click();", e)
                return True
    return False


def esperar_resultado_acta(driver, acta, timeout=10):
    lit = xpath_literal(acta)
    xp = f"//*[normalize-space(text())={lit}]"
    try:
        WebDriverWait(driver, timeout).until(lambda d: any(visible(e) for e in d.find_elements(By.XPATH, xp)))
        elems = [e for e in driver.find_elements(By.XPATH, xp) if visible(e)]
        return elems[-1] if elems else None
    except TimeoutException:
        return None


def clickable_from_element(el):
    try:
        tag = (el.tag_name or "").lower()
        if tag in {"a", "button", "input"}:
            return el
        if tag == "img":
            anc = el.find_elements(By.XPATH, "./ancestor::*[self::a or self::button][1]")
            return anc[0] if anc else el
    except Exception:
        pass
    return el


def click_tres_puntos(driver, acta_el):
    """Abre Información de Objeto desde la fila del resultado."""
    try:
        row = acta_el.find_elements(By.XPATH, "./ancestor::tr[1]")
        row = row[0] if row else None
    except Exception:
        row = None

    scope = row if row is not None else driver
    candidatos = []
    try:
        elems = scope.find_elements(By.XPATH, ".//*[self::a or self::button or self::input or self::img]" if row is not None else "//*[self::a or self::button or self::input or self::img]")
        seen = set()
        for raw in elems:
            e = clickable_from_element(raw)
            try:
                if not visible(e) or not e.is_enabled():
                    continue
                key = e.id
                if key in seen:
                    continue
                seen.add(key)
                r = e.rect
                if r.get("width", 0) <= 0 or r.get("height", 0) <= 0:
                    continue
                meta = " ".join([
                    e.text or "", e.get_attribute("title") or "", e.get_attribute("alt") or "",
                    e.get_attribute("aria-label") or "", e.get_attribute("class") or "",
                    e.get_attribute("src") or ""
                ])
                candidatos.append((e, norm(meta), r.get("x", 0)))
            except Exception:
                pass
    except Exception:
        pass

    # Preferencias semánticas.
    for palabra in ("DETALLE", "INFORMACION", "OBJETO", "ACCION", "OPCION", "MENU", "PUNTOS"):
        for e, meta, _ in candidatos:
            if palabra in meta:
                driver.execute_script("arguments[0].click();", e)
                return True
    for e, meta, _ in candidatos:
        if "..." in (e.text or "") or "•••" in (e.text or ""):
            driver.execute_script("arguments[0].click();", e)
            return True

    # En este sistema hay dos íconos al extremo derecho de la fila:
    # tres puntos y luego otro ícono. El botón de tres puntos es el segundo desde la derecha.
    if row is not None and len(candidatos) >= 2:
        candidatos.sort(key=lambda x: x[2])
        target = candidatos[-2][0]
        driver.execute_script("arguments[0].click();", target)
        return True
    return False


def _texto_contexto_actual(driver):
    try:
        return driver.find_element(By.TAG_NAME, "body").text or ""
    except Exception:
        return ""


def _buscar_contexto_con_texto(driver, frases, profundidad=0, max_profundidad=3):
    """Busca texto visible en documento principal o iframes y deja activo el contexto encontrado."""
    texto = norm(_texto_contexto_actual(driver))
    if all(norm(f) in texto for f in frases if f):
        return True

    if profundidad >= max_profundidad:
        return False

    try:
        frames = driver.find_elements(By.CSS_SELECTOR, "iframe, frame")
    except Exception:
        frames = []

    for fr in frames:
        try:
            if not visible(fr):
                continue
            driver.switch_to.frame(fr)
            if _buscar_contexto_con_texto(driver, frases, profundidad + 1, max_profundidad):
                return True
            driver.switch_to.parent_frame()
        except Exception:
            try:
                driver.switch_to.parent_frame()
            except Exception:
                pass
    return False


def _activar_contexto_info_objeto(driver):
    """Ubica el popup de Información de Objeto, incluso cuando GeneXus lo abre dentro de un iframe."""
    try:
        driver.switch_to.default_content()
    except Exception:
        pass

    # La captura del sistema muestra ambos textos; cualquiera de ellos identifica el popup.
    for frases in (["INFORMACION DE OBJETO"], ["ADMINISTRACION DE OBJETOS"], ["DATOS DESCRIPTIVOS", "INFRACCIONES"]):
        try:
            driver.switch_to.default_content()
        except Exception:
            pass
        if _buscar_contexto_con_texto(driver, frases):
            return True
    return False


def esperar_info_objeto(driver, timeout=10):
    fin = time.time() + timeout
    while time.time() < fin:
        if _activar_contexto_info_objeto(driver):
            return True
        time.sleep(0.25)
    return False


def click_lupa(driver, acta=None, log=None):
    """Abre Objetos Datos Descriptivos y confirma que la ventana realmente apareció.

    v6: el ícono amarillo de lupa del sistema GeneXus no siempre expone href/onclick
    en el IMG. Por eso se localiza la fila de datos y se prueban tanto el control
    real como puntos físicos dentro de la primera celda (donde está la lupa).
    """
    def _log(msg):
        if log:
            try:
                log(msg)
            except Exception:
                pass

    def _probar_y_confirmar(el, descripcion):
        try:
            driver.execute_script("arguments[0].scrollIntoView({block:'center',inline:'center'});", el)
            time.sleep(0.15)
            # click nativo primero; si GeneXus sólo escucha eventos del mouse, usar MouseEvent.
            try:
                el.click()
            except Exception:
                driver.execute_script("arguments[0].click();", el)
            if acta:
                texto = esperar_datos_descriptivos(driver, acta, timeout=2.2)
                if texto:
                    _log(f"  -> DATOS DESCRIPTIVOS abiertos ({descripcion})")
                    return texto
            else:
                return True
        except Exception:
            pass
        return None

    # Asegurar contexto del popup Información de Objeto.
    if not _activar_contexto_info_objeto(driver):
        return None if acta else False

    # 1) Localizar la tabla de Modelo / Marca / Color y su primera fila de datos.
    filas = []
    try:
        headers = driver.find_elements(
            By.XPATH,
            "//*[normalize-space(text())='Modelo' or normalize-space(text())='Marca' or normalize-space(text())='Color']"
        )
        for h in headers:
            if not visible(h):
                continue
            tables = h.find_elements(By.XPATH, "./ancestor::table[1]")
            for table in tables:
                rows = [r for r in table.find_elements(By.XPATH, ".//tr") if visible(r)]
                for row in rows[1:]:
                    txt = norm(row.text)
                    if txt and not ("MODELO" in txt and "MARCA" in txt and "COLOR" in txt):
                        filas.append(row)
                if filas:
                    break
            if filas:
                break
    except Exception:
        pass

    if filas:
        row = filas[0]
        _log("  -> Fila del vehículo detectada; buscando lupa real")

        # Primero: controles/elementos con onclick dentro de la primera celda.
        try:
            celdas = row.find_elements(By.XPATH, "./td")
            celda = celdas[0] if celdas else row
            candidatos = celda.find_elements(
                By.XPATH,
                ".//*[@onclick or @href or self::a or self::button or self::img or self::input or self::span]"
            )
            # incluir la propia celda por si el onclick vive allí
            candidatos.append(celda)
            vistos = set()
            for raw in candidatos:
                try:
                    if raw.id in vistos or not visible(raw):
                        continue
                    vistos.add(raw.id)
                    texto = _probar_y_confirmar(raw, "control primera celda")
                    if texto:
                        return texto
                    # tras un click que no abrió el detalle, volver a ubicar el popup
                    _activar_contexto_info_objeto(driver)
                except Exception:
                    pass
        except Exception:
            pass

        # Segundo: click físico dentro de la primera celda/fila.
        # En la captura la lupa está aproximadamente 15-25 px desde el borde izquierdo.
        try:
            target = (row.find_elements(By.XPATH, "./td") or [row])[0]
            ancho = max(1, int(target.rect.get("width", 40)))
            alto = max(1, int(target.rect.get("height", 25)))
            offsets = [8, 14, 20, 26, 32]
            for x in offsets:
                x = min(x, max(1, ancho - 2))
                try:
                    ActionChains(driver).move_to_element_with_offset(target, x - ancho/2, 0).click().perform()
                except Exception:
                    # alternativa con coordenadas absolutas del centro del elemento
                    try:
                        r = target.rect
                        cx = r.get("x", 0) + x
                        cy = r.get("y", 0) + alto/2
                        driver.execute_script(
                            "const e=document.elementFromPoint(arguments[0]-window.scrollX,arguments[1]-window.scrollY);"
                            "if(e){e.dispatchEvent(new MouseEvent('mousedown',{bubbles:true}));"
                            "e.dispatchEvent(new MouseEvent('mouseup',{bubbles:true}));e.click();}",
                            cx, cy
                        )
                    except Exception:
                        continue
                if acta:
                    texto = esperar_datos_descriptivos(driver, acta, timeout=2.2)
                    if texto:
                        _log(f"  -> DATOS DESCRIPTIVOS abiertos (click físico x={x})")
                        return texto
                else:
                    return True
                _activar_contexto_info_objeto(driver)
        except Exception:
            pass

    # 2) Fallback por geometría: imágenes/controles pequeños a la izquierda del modal.
    try:
        _activar_contexto_info_objeto(driver)
        elems = driver.find_elements(By.XPATH, "//*[self::img or self::a or self::button or self::span or @onclick]")
        candidatos = []
        for e in elems:
            if not visible(e):
                continue
            r = e.rect
            w, h = r.get("width", 0), r.get("height", 0)
            if 0 < w <= 60 and 0 < h <= 60:
                candidatos.append((e, r.get("x", 0), r.get("y", 0)))
        candidatos.sort(key=lambda z: (z[1], z[2]))
        for e, x, y in candidatos:
            meta = norm(" ".join([e.get_attribute("src") or "", e.get_attribute("title") or "", e.get_attribute("class") or ""]))
            if "CERR" in meta or "CLOSE" in meta:
                continue
            texto = _probar_y_confirmar(e, f"fallback x={round(x)} y={round(y)}")
            if texto:
                return texto
            _activar_contexto_info_objeto(driver)
    except Exception:
        pass

    return None if acta else False

def _recortar_bloque_datos_descriptivos(texto, acta=""):
    """Devuelve SOLO el bloque del modal final.

    GeneXus deja visible en el DOM el formulario de Consulta de Antecedentes detrás
    del modal. En v8 se leía el body completo y por eso podían capturarse valores
    como Tipo Vehículo=Todos (filtro de búsqueda) para todas las actas.
    """
    texto = str(texto or "")
    lineas = [x.rstrip() for x in texto.splitlines()]
    # Usar la ÚLTIMA aparición del título: el modal final suele estar al final del DOM.
    indices = [i for i, x in enumerate(lineas) if "OBJETOS DATOS DESCRIPTIVOS" in norm(x)]
    if indices:
        lineas = lineas[indices[-1]:]
    bloque = "\n".join(lineas).strip()

    # Si hubiera más de un bloque por ventanas anteriores, conservar el que contiene el acta actual.
    if acta and compact(acta) not in compact(bloque):
        # Buscar desde la última aparición del acta hacia atrás hasta el título más cercano.
        inds_acta = [i for i, x in enumerate(texto.splitlines()) if compact(acta) in compact(x)]
        if inds_acta:
            all_lines = texto.splitlines()
            ia = inds_acta[-1]
            inicio = 0
            for i in range(ia, -1, -1):
                if "OBJETOS DATOS DESCRIPTIVOS" in norm(all_lines[i]):
                    inicio = i
                    break
            bloque = "\n".join(all_lines[inicio:]).strip()
    return bloque


def _valor_control(el, driver):
    """Obtiene el valor real de INPUT/SELECT/TEXTAREA; .text no incluye INPUT.value."""
    try:
        tag = (el.tag_name or "").lower()
        if tag == "select":
            try:
                txt = driver.execute_script(
                    "return arguments[0].selectedOptions && arguments[0].selectedOptions.length ? arguments[0].selectedOptions[0].text : '';",
                    el,
                ) or ""
                if str(txt).strip():
                    return str(txt).strip()
            except Exception:
                pass
        val = el.get_attribute("value")
        if val is not None and str(val).strip():
            return str(val).strip()
        txt = el.text or ""
        return str(txt).strip()
    except Exception:
        return ""


def _contenedor_modal_final(driver):
    """Devuelve el contenedor visible más pequeño del popup OBJETOS DATOS DESCRIPTIVOS."""
    titulos = []
    try:
        titulos = driver.find_elements(
            By.XPATH,
            "//*[contains(translate(normalize-space(.),'abcdefghijklmnopqrstuvwxyzáéíóúñ','ABCDEFGHIJKLMNOPQRSTUVWXYZÁÉÍÓÚÑ'),'OBJETOS DATOS DESCRIPTIVOS')]"
        )
    except Exception:
        pass
    candidatos=[]
    for t in titulos:
        try:
            if not visible(t):
                continue
            cur=t
            for _ in range(10):
                try:
                    cur=cur.find_element(By.XPATH,"..")
                except Exception:
                    break
                tag=(cur.tag_name or '').lower()
                if tag in ('html','body'):
                    break
                txt=norm(cur.text or '')
                if 'OBJETOS DATOS DESCRIPTIVOS' not in txt:
                    continue
                # Debe contener varias etiquetas propias del formulario final.
                score=sum(1 for k in ('ACTA','TIPO VEHICULO','DOMINIO','MARCA','COLOR','MODELO') if k in txt)
                if score >= 4:
                    area=max(1, cur.rect.get('width',1))*max(1,cur.rect.get('height',1))
                    candidatos.append((area, -score, cur))
        except Exception:
            pass
    if candidatos:
        candidatos.sort(key=lambda x:(x[0],x[1]))
        return candidatos[0][2]
    return None


def _buscar_control_por_etiqueta(driver, contenedor, etiqueta):
    """Asocia una etiqueta visible con su input/select/textarea por DOM o proximidad visual."""
    etn=norm(etiqueta)
    labels=[]
    try:
        for e in contenedor.find_elements(By.XPATH, ".//*"):
            try:
                if not visible(e):
                    continue
                txt=(e.text or '').strip()
                if norm(txt)==etn:
                    labels.append(e)
            except Exception:
                pass
    except Exception:
        pass
    controles=[]
    try:
        controles=[e for e in contenedor.find_elements(By.CSS_SELECTOR,'input,select,textarea') if visible(e)]
    except Exception:
        pass
    invalid_types={'hidden','button','submit','image','checkbox','radio'}
    controles=[c for c in controles if (c.get_attribute('type') or '').lower() not in invalid_types]

    for lab in labels:
        # 1) label[for=id]
        try:
            fid=lab.get_attribute('for')
            if fid:
                for c in controles:
                    if c.get_attribute('id')==fid:
                        return c
        except Exception:
            pass
        # 2) misma celda / celda siguiente / fila
        xps=[
            "./ancestor::td[1]/following-sibling::td[1]//*[self::input or self::select or self::textarea]",
            "./ancestor::tr[1]//*[self::input or self::select or self::textarea]",
            "../following-sibling::*[1]//*[self::input or self::select or self::textarea]",
            "./following::*[self::input or self::select or self::textarea][1]",
        ]
        for xp in xps:
            try:
                cs=[c for c in lab.find_elements(By.XPATH,xp) if visible(c) and (c.get_attribute('type') or '').lower() not in invalid_types]
                if cs:
                    # si la fila tiene varios controles, elegir el más cercano al label
                    lr=lab.rect; lx=lr.get('x',0)+lr.get('width',0)/2; ly=lr.get('y',0)+lr.get('height',0)/2
                    cs.sort(key=lambda c: abs((c.rect.get('y',0)+c.rect.get('height',0)/2)-ly)*4 + abs((c.rect.get('x',0)+c.rect.get('width',0)/2)-lx))
                    return cs[0]
            except Exception:
                pass
        # 3) geometría: control a la derecha o justo debajo del label
        try:
            lr=lab.rect; lx=lr.get('x',0); ly=lr.get('y',0); lw=lr.get('width',0); lh=lr.get('height',0)
            scored=[]
            for c in controles:
                r=c.rect; cx=r.get('x',0); cy=r.get('y',0)
                dy=abs(cy-ly)
                # Preferir misma línea y a la derecha; permitir debajo.
                penalty=0 if cx >= lx-10 else 500
                score=dy*5 + abs(cx-(lx+lw)) + penalty
                if dy <= 80:
                    scored.append((score,c))
            if scored:
                scored.sort(key=lambda x:x[0])
                return scored[0][1]
        except Exception:
            pass
    return None


def _leer_modal_final_dom(driver, acta):
    """Lee valores REALES de los controles del popup final, no sólo body.text."""
    cont=_contenedor_modal_final(driver)
    if cont is None:
        return None
    etiquetas=['Acta','Fecha de Labrado','Tipo Vehículo','Dominio','Marca','Color','Modelo','Juzgado','COMISARIA']
    datos={}
    for et in etiquetas:
        ctrl=_buscar_control_por_etiqueta(driver,cont,et)
        val=_valor_control(ctrl,driver) if ctrl is not None else ''
        datos[et]=val

    # Acta puede estar renderizada como texto y no como input.
    if not datos.get('Acta'):
        txt=cont.text or ''
        m=re.search(r'(?im)^\s*Acta\s*[:\-]?\s*([^\r\n]+)$',txt)
        if m:
            datos['Acta']=m.group(1).strip()
    # No aceptar una lectura claramente ajena al acta actual.
    if datos.get('Acta') and compact(acta) not in compact(datos.get('Acta')):
        return None

    # Deben existir al menos dos datos descriptivos reales; evita falsos AUTOMOVIL + vacíos.
    llenos=sum(bool(str(datos.get(k,'')).strip()) for k in ('Tipo Vehículo','Marca','Color','Dominio','Modelo'))
    if llenos < 2:
        return None
    return "\n".join(f"{k} {v}" for k,v in datos.items() if str(v).strip())


def esperar_datos_descriptivos(driver, acta, timeout=10):
    """Espera el popup final y extrae valores de sus INPUT/SELECT reales."""
    fin=time.time()+timeout
    while time.time()<fin:
        try:
            driver.switch_to.default_content()
        except Exception:
            pass
        if _buscar_contexto_con_texto(driver,['OBJETOS DATOS DESCRIPTIVOS']):
            texto=_leer_modal_final_dom(driver,acta)
            if texto:
                return texto
        time.sleep(0.25)
    return None

def extraer_campo(texto, etiqueta, siguientes):
    """Extrae la ÚLTIMA ocurrencia válida de una etiqueta dentro del modal final."""
    lines = [re.sub(r"\s+", " ", x).strip() for x in str(texto or "").splitlines() if x.strip()]
    et = norm(etiqueta)
    sigs = {norm(s) for s in siguientes}
    invalidos = {"TODOS", "TODAS", "SELECCIONE", "SELECCIONAR", "--", "-"}
    candidatos = []

    for i, line in enumerate(lines):
        nl = norm(line)
        if nl == et or nl.startswith(et + " "):
            # etiqueta + valor en la misma línea
            m = re.match(rf"(?i)^\s*{re.escape(etiqueta)}\s*[:\-]?\s*(.+)$", line)
            if m:
                val = m.group(1).strip()
                if val and norm(val) not in invalidos and norm(val) != et:
                    candidatos.append(val)
                    continue
            # etiqueta y valor en línea siguiente; saltar líneas vacías/etiquetas
            for j in range(i + 1, min(len(lines), i + 6)):
                nv = norm(lines[j])
                if nv in sigs:
                    break
                if not nv or nv in invalidos:
                    continue
                candidatos.append(lines[j])
                break

    return candidatos[-1] if candidatos else ""


def extraer_datos(texto):
    etiquetas = [
        "Acta", "Fecha de Labrado", "Retiene Licencia", "Número de Licencia",
        "Categoria Licencia", "Procedencia Licencia", "Vto. de Licencia",
        "Tipo Vehículo", "Dominio", "Marca", "Color", "Modelo", "Juzgado",
        "COMISARIA", "Usuario", "Lugar"
    ]
    datos = {}
    for e in etiquetas:
        datos[e] = extraer_campo(texto, e, etiquetas)

    # Fallback regex multiline específico para los campos que necesitamos.
    patrones = {
        "Acta": r"(?im)^\s*Acta\s+([^\r\n]+)$",
        "Fecha de Labrado": r"(?im)^\s*Fecha\s+de\s+Labrado\s+([^\r\n]+)$",
        "Tipo Vehículo": r"(?im)^\s*Tipo\s+Veh[ií]culo\s+([^\r\n]+)$",
        "Dominio": r"(?im)^\s*Dominio\s+([^\r\n]+)$",
        "Marca": r"(?im)^\s*Marca\s+([^\r\n]+)$",
        "Color": r"(?im)^\s*Color\s+([^\r\n]+)$",
        "Modelo": r"(?im)^\s*Modelo\s+([^\r\n]+)$",
        "Juzgado": r"(?im)^\s*Juzgado\s+([^\r\n]+)$",
        "COMISARIA": r"(?im)^\s*COMISARIA\s+([^\r\n]+)$",
    }
    for k, pat in patrones.items():
        if not datos.get(k):
            ms = list(re.finditer(pat, texto or ""))
            if ms:
                datos[k] = ms[-1].group(1).strip()

    # Nunca aceptar valores propios de los filtros de la pantalla principal.
    for k in ("Tipo Vehículo", "Marca", "Color", "Modelo", "Dominio"):
        if norm(datos.get(k, "")) in {"TODOS", "TODAS", "SELECCIONE", "SELECCIONAR"}:
            datos[k] = ""
    return datos


def extraer_primera_acta_vial_de_texto(texto):
    """Devuelve la primera Acta Vial L/X/R/F visible, en el orden del resultado."""
    texto = str(texto or "").upper()
    patron = re.compile(r"(?<![A-Z0-9])([LXRF])\s*[- ]?\s*(\d{5,12})(?!\d)")
    for m in patron.finditer(texto):
        acta = acta_canonica(m.group(1) + m.group(2))
        if es_acta_vial(acta):
            return acta
    return None


def asegurar_columna_acta_sugerida(ws, header_row):
    """Crea/reutiliza una columna separada; nunca pisa la celda original de Acta/Sumario."""
    for c in range(1, ws.max_column + 1):
        if norm(ws.cell(header_row, c).value) in {"ACTA VIAL SUGERIDA", "ACTA VIAL DETECTADA"}:
            return c
    c = ws.max_column + 1
    ws.cell(header_row, c).value = "ACTA VIAL SUGERIDA"
    return c


def ruta_recurso(nombre):
    """Devuelve la ruta de un recurso tanto en .py como dentro del EXE de PyInstaller."""
    candidatos=[]
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        candidatos.append(os.path.join(sys._MEIPASS, nombre))
    candidatos.append(os.path.join(directorio_app(), nombre))
    for r in candidatos:
        if os.path.exists(r):
            return r
    return candidatos[-1]


def obtener_modelo_resolucion_final():
    """Obtiene el modelo de Disposición Final de forma robusta.

    Orden de búsqueda:
    1) recurso de PyInstaller (_MEIPASS),
    2) archivo junto al programa,
    3) copia embebida dentro del propio EXE.
    """
    ruta = ruta_recurso("modelo_resolucion_final.xlsx")
    if ruta and os.path.exists(ruta):
        return ruta

    try:
        from modelo_resolucion_embebido import obtener_modelo_bytes
        import tempfile
        carpeta = os.path.join(tempfile.gettempdir(), "Verificador_Vial_Caminera")
        os.makedirs(carpeta, exist_ok=True)
        destino = os.path.join(carpeta, "modelo_resolucion_final_v20.xlsx")
        datos = obtener_modelo_bytes()
        # Regenerar si no existe o quedó incompleto.
        if (not os.path.exists(destino)) or os.path.getsize(destino) != len(datos):
            with open(destino, "wb") as f:
                f.write(datos)
        return destino
    except Exception as e:
        raise RuntimeError(
            "No pude cargar el modelo de Disposición Final. "
            f"Detalle: {e}"
        )


def buscar_columna_por_encabezado(ws, header_row, grupos):
    """Busca una columna por sinónimos de encabezado. grupos = [(palabras obligatorias), ...]."""
    for c in range(1, ws.max_column + 1):
        h = _header(ws.cell(header_row, c).value)
        if not h:
            continue
        for grupo in grupos:
            if all(p in h for p in grupo):
                return c
    return None


def detectar_columna_juzgado(ws, header_row, cfg=None):
    if cfg and cfg.get("interviene"):
        return cfg["interviene"]
    return buscar_columna_por_encabezado(ws, header_row, [
        ("JUZGADO",), ("INTERVIENE",), ("INTERVINIENTE",), ("JDO",)
    ])


def normalizar_juzgado_para_titulo(v):
    s = str(v or "").strip()
    if not s:
        return ""
    # Mantener la denominación original, sólo compactar espacios y saltos.
    return re.sub(r"\s+", " ", s).strip()


def juzgados_presentes(ws, filas, header_row, cfg=None):
    col = detectar_columna_juzgado(ws, header_row, cfg)
    if not col:
        return []
    vistos=set(); salida=[]
    for fila in filas:
        j = normalizar_juzgado_para_titulo(ws.cell(fila, col).value)
        if not j:
            continue
        clave = norm(j)
        if clave in {"NO CONSTA", "SIN DATO", "S D", "SD", "NINGUNO", "-"}:
            continue
        if clave not in vistos:
            vistos.add(clave)
            salida.append(j)
    return salida


def _copiar_estilo_celda(origen, destino):
    try:
        from copy import copy
        if origen.has_style:
            destino._style = copy(origen._style)
        destino.font = copy(origen.font)
        destino.fill = copy(origen.fill)
        destino.border = copy(origen.border)
        destino.alignment = copy(origen.alignment)
        destino.number_format = origen.number_format
        destino.protection = copy(origen.protection)
    except Exception:
        pass


def _preparar_hoja_disposicion(ws, nombre_hoja, subtitulo, juzgados, filas_datos):
    """Limpia el modelo dejando títulos/encabezados y carga los datos finales."""
    ws.title = nombre_hoja[:31]
    # Guardar estilo de la primera fila de datos del modelo antes de borrar.
    estilos=[]
    valores_modelo=[]
    for c in range(1, 9):
        cel=ws.cell(4,c)
        estilos.append(cel)
        valores_modelo.append(cel.value)
    altura_modelo = ws.row_dimensions[4].height

    if ws.max_row >= 4:
        ws.delete_rows(4, ws.max_row - 3)

    # Título: mostrar los juzgados presentes AL LADO DE PRO.ME.COM.
    texto_juzgados = " - JUZGADOS: " + " - ".join(juzgados) if juzgados else " - JUZGADOS: SIN DATO"
    ws["E1"] = "PRO.ME.COM" + texto_juzgados
    try:
        from copy import copy
        al=copy(ws["E1"].alignment); al.wrap_text=True; al.horizontal="center"; al.vertical="center"; ws["E1"].alignment=al
    except Exception:
        pass
    try:
        # Tamaño dinámico para títulos largos.
        from copy import copy
        largo=len(ws["E1"].value or "")
        tam=11 if largo <= 90 else 9 if largo <= 150 else 8
        ft=copy(ws["E1"].font); ft.sz=tam; ft.bold=True; ws["E1"].font=ft
        ws.row_dimensions[1].height = 28 if largo <= 110 else 42
    except Exception:
        pass

    # Conservar el título territorial del modelo y agregar el tipo de listado.
    titulo_base = str(ws["E2"].value or "").strip()
    if subtitulo:
        if titulo_base:
            ws["E2"] = f"{titulo_base}    |    {subtitulo}"
        else:
            ws["E2"] = subtitulo
    try:
        from copy import copy
        al=copy(ws["E2"].alignment); al.wrap_text=True; al.horizontal="center"; al.vertical="center"; ws["E2"].alignment=al
    except Exception:
        pass

    for idx, datos in enumerate(filas_datos, start=1):
        r = idx + 3
        fila_val=[idx] + list(datos)
        for c, val in enumerate(fila_val, start=1):
            dst=ws.cell(r,c)
            dst.value=val
            _copiar_estilo_celda(estilos[c-1], dst)
        if altura_modelo:
            ws.row_dimensions[r].height = altura_modelo

    # Repetir encabezados al imprimir y ajustar área.
    try:
        ws.print_title_rows = '1:3'
        ws.print_area = f'A1:H{max(3, len(filas_datos)+3)}'
    except Exception:
        pass


def generar_disposicion_final_desde_archivo(ruta_origen, ruta_salida=None):
    """Genera un Excel de Disposición Final con VERIFICADOS y CON DIFERENCIAS.

    Usa el modelo oficial suministrado por el usuario y muestra en E1, al lado de
    PRO.ME.COM, los juzgados que efectivamente aparecen en cada listado.
    """
    if not ruta_origen or not os.path.exists(ruta_origen):
        raise RuntimeError("No se encontró el Excel verificado para generar la Disposición Final.")

    wb_src = load_workbook(ruta_origen, data_only=False)
    cfg = detectar_estructura_excel(wb_src)
    ws_src = cfg["sheet"]
    h = cfg["header_row"]

    col_interno = buscar_columna_por_encabezado(ws_src, h, [
        ("INTERNO",), ("REGISTRO",), ("NRO", "REG"), ("N°", "REG")
    ])
    col_motor = buscar_columna_por_encabezado(ws_src, h, [("MOTOR",)])
    col_acta_sug = buscar_columna_por_encabezado(ws_src, h, [("ACTA", "VIAL", "SUGERIDA"), ("ACTA", "VIAL", "DETECTADA")])

    verificadas=[]
    diferencias=[]
    filas_ver=[]
    filas_dif=[]

    for fila in range(h+1, ws_src.max_row+1):
        estado = str(ws_src.cell(fila, cfg["verificacion"]).value or "").strip()
        ne = norm(estado)
        if not ne:
            continue
        es_ver = "VERIFICADO" in ne and "NO VERIFICADO" not in ne
        es_dif = ne.startswith("NO COINCIDE") or "DIFERENCIAS" in ne
        if not (es_ver or es_dif):
            continue

        interno = ws_src.cell(fila, col_interno).value if col_interno else ""
        tipo = ws_src.cell(fila, cfg["tipo"]).value if cfg.get("tipo") else ""
        marca = ws_src.cell(fila, cfg["marca_modelo"]).value if cfg.get("marca_modelo") else ""
        dominio = ws_src.cell(fila, cfg["dominio"]).value if cfg.get("dominio") else ""
        color = ws_src.cell(fila, cfg["color"]).value if cfg.get("color") else ""
        motor = ws_src.cell(fila, col_motor).value if col_motor else ""
        sumario = ws_src.cell(fila, cfg["acta"]).value if cfg.get("acta") else ""
        if (sumario is None or not str(sumario).strip()) and col_acta_sug:
            sug=ws_src.cell(fila, col_acta_sug).value
            if sug and not str(sug).upper().startswith("ERROR") and "NO SE ENCONTRO" not in norm(sug):
                sumario=sug
        datos=[interno or "", tipo or "", marca or "", dominio or "", color or "", motor or "", sumario or ""]
        if es_ver:
            verificadas.append(datos); filas_ver.append(fila)
        else:
            diferencias.append(datos); filas_dif.append(fila)

    modelo = obtener_modelo_resolucion_final()
    wb_out = load_workbook(modelo)
    base = wb_out[wb_out.sheetnames[0]]
    hoja_dif = wb_out.copy_worksheet(base)

    j_ver = juzgados_presentes(ws_src, filas_ver, h, cfg)
    j_dif = juzgados_presentes(ws_src, filas_dif, h, cfg)

    _preparar_hoja_disposicion(base, "VERIFICADOS", "VEHÍCULOS VERIFICADOS", j_ver, verificadas)
    _preparar_hoja_disposicion(hoja_dif, "CON DIFERENCIAS", "VEHÍCULOS CON DIFERENCIAS", j_dif, diferencias)

    if ruta_salida is None:
        carpeta,nombre=os.path.split(ruta_origen)
        base_nombre=os.path.splitext(nombre)[0]
        ruta_salida=os.path.join(carpeta, f"{base_nombre}_DISPOSICION_FINAL.xlsx")
    wb_out.save(ruta_salida)
    return ruta_salida, len(verificadas), len(diferencias), j_ver, j_dif


class App:
    def __init__(self, root):
        self.root = root
        self.root.title(f"Verificador Vial Caminera v{VERSION}")
        self.root.geometry("1040x740")
        self.root.minsize(900, 650)
        self.driver = None
        self.consulta_url = None
        self.stop_flag = False
        self.ultimo_archivo_salida = None
        self.archivo = tk.StringVar()
        self.solo_pendientes = tk.BooleanVar(value=True)
        self.limite = tk.StringVar(value="1")
        self._ui()

    def _ui(self):
        tk.Label(self.root, text="VERIFICADOR VIAL CAMINERA", font=("Segoe UI", 18, "bold")).pack(pady=(15, 3))
        tk.Label(self.root, text=f"Versión {VERSION} UNIVERSAL · Actas L/X/R/F · dominio normal/invertido · Disposición Final · modelo embebido.", font=("Segoe UI", 10)).pack(pady=(0, 4))
        tk.Label(self.root, text="No guarda usuario ni contraseña. La sesión se inicia manualmente en Chrome.", font=("Segoe UI", 9)).pack(pady=(0, 12))

        f = tk.Frame(self.root)
        f.pack(fill="x", padx=20)
        tk.Entry(f, textvariable=self.archivo, font=("Segoe UI", 10)).pack(side="left", fill="x", expand=True)
        tk.Button(f, text="Elegir Excel", command=self.elegir).pack(side="left", padx=(8, 0))

        opts = tk.Frame(self.root)
        opts.pack(fill="x", padx=20, pady=10)
        tk.Checkbutton(opts, text="Saltar filas ya VERIFICADAS", variable=self.solo_pendientes).pack(side="left")
        tk.Label(opts, text="Máx. filas (0 = todas):").pack(side="left", padx=(25, 0))
        tk.Entry(opts, textvariable=self.limite, width=5).pack(side="left", padx=5)
        tk.Label(opts, text="Para la primera prueba deje 1.", font=("Segoe UI", 9, "italic")).pack(side="left", padx=8)

        botones = tk.Frame(self.root)
        botones.pack(fill="x", padx=20, pady=(0, 10))
        tk.Button(botones, text="1. ABRIR SISTEMA / INICIAR SESIÓN", command=self.abrir_sistema, height=2).pack(side="left", fill="x", expand=True)
        tk.Button(botones, text="2. INICIAR VERIFICACIÓN", command=self.iniciar, height=2).pack(side="left", fill="x", expand=True, padx=8)
        tk.Button(botones, text="3. GENERAR DISPOSICIÓN FINAL", command=self.generar_disposicion, height=2).pack(side="left", fill="x", expand=True)
        tk.Button(botones, text="DETENER", command=self.detener, height=2).pack(side="left", padx=(8,0))

        self.estado = tk.Label(self.root, text="Listo.", anchor="w", font=("Segoe UI", 10, "bold"))
        self.estado.pack(fill="x", padx=20, pady=(2, 5))
        self.logbox = ScrolledText(self.root, height=26, font=("Consolas", 9))
        self.logbox.pack(fill="both", expand=True, padx=20, pady=(0, 15))

    def log(self, msg):
        ts = datetime.now().strftime("%H:%M:%S")
        self.root.after(0, lambda: (self.logbox.insert("end", f"[{ts}] {msg}\n"), self.logbox.see("end")))

    def set_estado(self, msg):
        self.root.after(0, lambda: self.estado.config(text=msg))

    def elegir(self):
        p = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if p:
            self.archivo.set(p)
            self.log(f"Excel seleccionado: {p}")

    def generar_disposicion(self):
        """Genera el Excel final sin necesidad de tener Chrome abierto."""
        origen = self.ultimo_archivo_salida
        if not origen or not os.path.exists(origen):
            p=self.archivo.get()
            if p and os.path.exists(p):
                carpeta,nombre=os.path.split(p)
                base,ext=os.path.splitext(nombre)
                candidato=os.path.join(carpeta, f"{base}_VERIFICADO{ext}")
                origen = candidato if os.path.exists(candidato) else p
        if not origen or not os.path.exists(origen):
            messagebox.showwarning("Disposición Final", "Seleccione primero el Excel verificado.")
            return
        try:
            self.set_estado("Generando Disposición Final...")
            salida,nver,ndif,jver,jdif=generar_disposicion_final_desde_archivo(origen)
            self.log(f"Disposición Final generada: {salida}")
            self.log(f"  VERIFICADOS: {nver} | Juzgados presentes: {', '.join(jver) if jver else 'SIN DATO'}")
            self.log(f"  CON DIFERENCIAS: {ndif} | Juzgados presentes: {', '.join(jdif) if jdif else 'SIN DATO'}")
            self.set_estado("Disposición Final generada.")
            messagebox.showinfo("Disposición Final", f"Archivo generado correctamente.\n\nVerificados: {nver}\nCon diferencias: {ndif}\n\n{salida}")
        except Exception as e:
            self.set_estado("Error generando Disposición Final.")
            self.log(f"ERROR GENERANDO DISPOSICIÓN FINAL: {e}")
            messagebox.showerror("Disposición Final", str(e))

    def abrir_sistema(self):
        try:
            if self.driver:
                self.driver.quit()
        except Exception:
            pass
        try:
            opts = webdriver.ChromeOptions()
            opts.add_argument("--start-maximized")
            opts.add_experimental_option("detach", True)
            self.driver = webdriver.Chrome(options=opts)
            self.driver.get(URL_LOGIN)
            self.consulta_url = None
            self.log("Chrome abierto. Inicie sesión normalmente y entre a CONSULTA DE ANTECEDENTES.")
            self.set_estado("Esperando inicio de sesión...")
            messagebox.showinfo("Paso 1", "Inicie sesión en Chrome y abra 'CONSULTA DE ANTECEDENTES'.\n\nDespués vuelva al programa y pulse 'INICIAR VERIFICACIÓN'.")
        except WebDriverException as e:
            messagebox.showerror("Chrome/Selenium", f"No pude abrir Chrome.\n\n{e}")

    def detener(self):
        self.stop_flag = True
        self.log("Se solicitó detener el proceso.")

    def iniciar(self):
        if not self.archivo.get() or not os.path.exists(self.archivo.get()):
            messagebox.showwarning("Excel", "Seleccione primero el archivo Excel.")
            return
        if not self.driver:
            messagebox.showwarning("Sistema", "Primero pulse 'ABRIR SISTEMA / INICIAR SESIÓN'.")
            return
        try:
            if URL_CONSULTA_HINT not in (self.driver.current_url or "").lower():
                messagebox.showwarning("Sistema", "Chrome todavía no está en 'Consulta de Antecedentes'.\n\nAbra esa opción en el sistema y vuelva a intentar.")
                return
            self.consulta_url = self.driver.current_url
        except Exception:
            pass
        self.stop_flag = False
        threading.Thread(target=self.procesar, daemon=True).start()

    def volver_consulta(self):
        if not self.consulta_url:
            return
        self.driver.get(self.consulta_url)
        WebDriverWait(self.driver, 8).until(lambda d: find_input_acta(d) is not None)

    def consultar_acta(self, acta):
        """Busca el acta y abre sólo Información de Objeto.

        La versión universal elimina la lupa/Datos Descriptivos. Para vehículos sin dominio la
        identificación se hace con Marca/Modelo + Color de esta primera ventana.
        """
        self.volver_consulta()
        inp = find_input_acta(self.driver, self.log)
        if not inp:
            raise RuntimeError("No encontré el campo Nro Acta.")
        cargado=False
        for intento in range(1,4):
            inp=find_input_acta(self.driver, self.log)
            if not inp:
                return None, "NO SE ENCONTRO CAMPO NRO ACTA"
            try:
                if cargar_input_geneXus(self.driver, inp, str(acta)):
                    valor=(inp.get_attribute("value") or "").strip()
                    self.log(f"  Campo Nro Acta intento {intento}: '{valor}'")
                    if compact(valor)==compact(acta):
                        cargado=True
                        break
            except Exception as ex:
                self.log(f"  Reintento de carga del acta: {ex}")
            time.sleep(0.4)
        if not cargado:
            return None, "EL SISTEMA NO ACEPTO EL NRO ACTA"

        if not click_buscar(self.driver):
            inp.send_keys(Keys.ENTER)
        time.sleep(0.5)
        acta_el=esperar_resultado_acta(self.driver, acta, timeout=10)
        if not acta_el:
            return None, "NO ENCONTRADA"
        if not click_tres_puntos(self.driver, acta_el):
            return None, "NO SE PUDO ABRIR LOS TRES PUNTOS"
        if not esperar_info_objeto(self.driver, timeout=8):
            return None, "NO SE DETECTO INFORMACION DE OBJETO"
        self.log("  -> INFORMACION DE OBJETO detectada")
        datos=extraer_info_objeto_simple(self.driver)
        return datos, None

    def consultar_dominio_contiene_acta(self, dominio, actas):
        dominios = candidatos_dominio_verificacion(dominio)
        if not dominios:
            return None, None, None
        candidatos = actas if isinstance(actas, (list, tuple)) else [actas]
        candidatos = [a for a in candidatos if a]

        for idx, dominio_prueba in enumerate(dominios, start=1):
            self.volver_consulta()
            inp=find_input_dominio(self.driver, self.log)
            if not inp:
                return False, "NO SE ENCONTRO CAMPO DOMINIO", None
            if not cargar_input_geneXus(self.driver, inp, dominio_prueba):
                return False, "EL SISTEMA NO ACEPTO EL DOMINIO", None
            if len(dominios) > 1:
                self.log(f"  Verificación principal: Dominio {dominio_prueba} (variante {idx}/{len(dominios)}) -> Acta(s) {', '.join(candidatos)}")
            else:
                self.log(f"  Verificación principal: Dominio {dominio_prueba} -> Acta(s) {', '.join(candidatos)}")
            if not click_buscar(self.driver):
                inp.send_keys(Keys.ENTER)
            time.sleep(0.7)
            try:
                WebDriverWait(self.driver, 10).until(
                    lambda d: any(cuerpo_contiene_acta_equivalente(d.find_element(By.TAG_NAME,'body').text, a) for a in candidatos)
                    or 'NO SE ENCONTR' in norm(d.find_element(By.TAG_NAME,'body').text)
                )
            except Exception:
                pass
            texto=self.driver.find_element(By.TAG_NAME,'body').text
            for a in candidatos:
                if cuerpo_contiene_acta_equivalente(texto, a):
                    if idx > 1:
                        self.log(f"  -> Coincidencia hallada usando dominio invertido: {dominio_prueba}")
                    return True, None, a
        return False, None, None

    def consultar_dominio_primera_acta(self, dominio):
        """Busca un dominio y devuelve la primera Acta Vial L/X/R/F visible.

        Se usa EXCLUSIVAMENTE cuando la celda original destinada al Acta Vial está vacía.
        Prueba también la orientación invertida de dominios clásicos cuando corresponde.
        """
        dominios = candidatos_dominio_verificacion(dominio)
        if not dominios:
            return None, "DOMINIO NO UTILIZABLE", None
        for idx, dominio_prueba in enumerate(dominios, start=1):
            self.volver_consulta()
            inp = find_input_dominio(self.driver, self.log)
            if not inp:
                return None, "NO SE ENCONTRO CAMPO DOMINIO", None
            if not cargar_input_geneXus(self.driver, inp, dominio_prueba):
                return None, "EL SISTEMA NO ACEPTO EL DOMINIO", None
            self.log(f"  Celda Acta Vial VACÍA: buscando primera acta para dominio {dominio_prueba} ({idx}/{len(dominios)})")
            if not click_buscar(self.driver):
                inp.send_keys(Keys.ENTER)
            time.sleep(0.8)
            try:
                WebDriverWait(self.driver, 10).until(
                    lambda d: extraer_primera_acta_vial_de_texto(d.find_element(By.TAG_NAME,'body').text) is not None
                    or 'NO SE ENCONTR' in norm(d.find_element(By.TAG_NAME,'body').text)
                )
            except Exception:
                pass
            texto = self.driver.find_element(By.TAG_NAME,'body').text
            acta = extraer_primera_acta_vial_de_texto(texto)
            if acta:
                return acta, None, dominio_prueba
        return None, None, None

    def procesar(self):
        try:
            limite=int(self.limite.get() or "0")
        except ValueError:
            limite=0

        p=self.archivo.get()
        carpeta,nombre=os.path.split(p)
        base,ext=os.path.splitext(nombre)
        stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
        respaldo=os.path.join(carpeta,f"{base}_RESPALDO_{stamp}{ext}")
        salida=os.path.join(carpeta,f"{base}_VERIFICADO{ext}")
        shutil.copy2(p,respaldo)
        self.log(f"Respaldo creado: {os.path.basename(respaldo)}")

        wb=load_workbook(p)
        try:
            cfg=detectar_estructura_excel(wb)
        except Exception as e:
            self.log(f"ERROR DETECTANDO ESTRUCTURA: {e}")
            self.root.after(0,lambda: messagebox.showerror("Estructura Excel", str(e)))
            return
        ws=cfg["sheet"]
        col_acta_sugerida = asegurar_columna_acta_sugerida(ws, cfg["header_row"])
        total=verificados=diferencias=no_encontrados=errores=0

        def colname(n):
            try: return ws.cell(cfg["header_row"],n).column_letter
            except Exception: return "?"
        self.log(
            f"Estructura detectada: hoja='{cfg['sheet_name']}' | encabezados fila {cfg['header_row']} | "
            f"Tipo={colname(cfg.get('tipo')) if cfg.get('tipo') else '-'} | "
            f"Marca/Modelo={colname(cfg['marca_modelo'])} | Color={colname(cfg['color'])} | "
            f"Dominio={colname(cfg['dominio'])} | Acta={colname(cfg['acta'])} | "
            f"Verificación={colname(cfg['verificacion'])}"
        )
        # Avisar si COLOR/DOMINIO quedaron intercambiados respecto del rótulo: ocurre en San Cristóbal.
        h_color=norm(ws.cell(cfg['header_row'],cfg['color']).value)
        h_dom=norm(ws.cell(cfg['header_row'],cfg['dominio']).value)
        if "DOMINIO" in h_color or "COLOR" in h_dom:
            self.log("  -> Se detectó encabezado COLOR/DOMINIO desplazado; se usarán los datos reales de las columnas.")

        for fila in range(cfg["header_row"]+1, ws.max_row+1):
            if self.stop_flag: break
            obs=str(ws.cell(fila,cfg["verificacion"]).value or "").strip().upper()
            if self.solo_pendientes.get() and "VERIFICADO" in obs:
                continue

            acta, acta_original = extraer_acta_de_fila(ws,fila,cfg)
            acta_celda_original = ws.cell(fila, cfg["acta"]).value if cfg.get("acta") else None
            acta_celda_vacia = acta_celda_original is None or not str(acta_celda_original).strip()
            dominio_excel=ws.cell(fila,cfg["dominio"]).value
            dominio_real=extraer_dominio_real(dominio_excel)
            dominios_prueba=candidatos_dominio_verificacion(dominio_excel)
            marca_excel=ws.cell(fila,cfg["marca_modelo"]).value
            color_excel=ws.cell(fila,cfg["color"]).value

            # Ignorar filas totalmente vacías / títulos / separadores.
            if not any(str(x or "").strip() for x in (acta_original,dominio_excel,marca_excel,color_excel)):
                continue
            if limite and total>=limite: break
            total+=1

            referencia=acta or str(acta_original or "SIN ACTA VIAL").strip()
            self.set_estado(f"Fila {fila} · {referencia}")
            self.log(f"Consultando fila {fila} - Acta {referencia}...")
            if acta and acta_original is not None and compact(str(acta_original)) != compact(acta):
                self.log(f"  Acta normalizada: '{acta_original}' -> '{acta}'")
            self.log(f"  Excel detectado: Marca/Modelo={marca_excel or ''} | Color={color_excel or ''} | Dominio={dominio_excel or ''}")
            if len(dominios_prueba) > 1:
                self.log(f"  Dominio clásico: se probará también invertido -> {', '.join(dominios_prueba)}")
            candidatos_acta = candidatos_acta_verificacion(acta, acta_original)
            if len(candidatos_acta) > 1:
                self.log(f"  Acta sin letra: se probará con {', '.join(candidatos_acta)}")

            try:
                # REGLA ESPECIAL v18: SOLO si la CELDA ORIGINAL de Acta Vial está VACÍA
                # y existe un dominio real, buscar qué primera Acta Vial L/X/R/F figura asociada.
                # No modifica la celda original: escribe el hallazgo en ACTA VIAL SUGERIDA.
                if dominio_real and acta_celda_vacia:
                    sugerida, error_sug, dominio_usado = self.consultar_dominio_primera_acta(str(dominio_excel or ''))
                    if error_sug:
                        ws.cell(fila, col_acta_sugerida).value = f"ERROR: {error_sug}"
                        self.log(f"  -> No se pudo determinar Acta Vial sugerida: {error_sug}")
                    elif sugerida:
                        ws.cell(fila, col_acta_sugerida).value = sugerida
                        self.log(f"  -> ACTA VIAL SUGERIDA: {sugerida} (dominio consultado: {dominio_usado})")
                    else:
                        ws.cell(fila, col_acta_sugerida).value = "NO SE ENCONTRO ACTA VIAL ASOCIADA"
                        self.log("  -> No se encontró Acta Vial L/X/R/F asociada al dominio")
                    # Esta regla sólo genera la sugerencia. No transforma la fila en VERIFICADA.
                    wb.save(salida)
                    continue

                # REGLA 1: CON DOMINIO REAL, el cruce DOMINIO -> ACTA es la prueba principal.
                if dominio_real:
                    if not acta:
                        ws.cell(fila,cfg["verificacion"]).value="NO VERIFICABLE: DOMINIO PRESENTE PERO SIN ACTA VIAL EN EXCEL"
                        no_encontrados+=1
                        self.log(f"  -> Dominio real {dominio_real}, pero no hay Acta Vial utilizable en la fila")
                        wb.save(salida)
                        continue
                    cruce_ok,cruce_error,acta_match=self.consultar_dominio_contiene_acta(str(dominio_excel or ''),candidatos_acta or [acta])
                    if cruce_error:
                        ws.cell(fila,cfg["verificacion"]).value=f"ERROR VERIFICANDO DOMINIO: {cruce_error}"
                        errores+=1
                        self.log(f"  -> {cruce_error}")
                    elif cruce_ok:
                        ws.cell(fila,cfg["verificacion"]).value="VERIFICADO POR DOMINIO"
                        verificados+=1
                        self.log(f"  -> VERIFICADO POR DOMINIO ({dominio_real} ↔ {acta_match or acta})")
                    else:
                        ws.cell(fila,cfg["verificacion"]).value="NO COINCIDE: DOMINIO NO ASOCIA ACTA"
                        diferencias+=1
                        self.log(f"  -> DOMINIO {dominio_real} NO ASOCIA NINGUNA ACTA PROBADA: {', '.join(candidatos_acta or [acta])}")
                    wb.save(salida)
                    continue

                # REGLA 2: SIN DOMINIO REAL, necesitamos Acta Vial + Información de Objeto.
                if not acta:
                    ws.cell(fila,cfg["verificacion"]).value="NO VERIFICABLE: SIN DOMINIO NI ACTA VIAL"
                    no_encontrados+=1
                    self.log("  -> NO VERIFICABLE: SIN DOMINIO NI ACTA VIAL")
                    wb.save(salida)
                    continue

                datos=None
                problema="NO ENCONTRADA"
                acta_usada=acta
                for candidata in (candidatos_acta or [acta]):
                    if len(candidatos_acta or []) > 1:
                        self.log(f"  Probando acta sin dominio como {candidata}...")
                    datos,problema=self.consultar_acta(candidata)
                    if not problema:
                        acta_usada=candidata
                        if candidata != acta:
                            self.log(f"  -> Acta encontrada utilizando prefijo: {candidata}")
                        break
                    if problema != "NO ENCONTRADA":
                        break
                if problema:
                    if problema=="NO ENCONTRADA":
                        if es_acta_vial(acta):
                            ws.cell(fila,cfg["verificacion"]).value="NO ENCONTRADA EN SISTEMA VIAL"
                            no_encontrados+=1
                            self.log("  -> NO ENCONTRADA EN SISTEMA VIAL")
                        else:
                            ws.cell(fila,cfg["verificacion"]).value="NO VERIFICABLE POR ACTA VIAL"
                            no_encontrados+=1
                            self.log("  -> NO VERIFICABLE POR ACTA VIAL")
                    else:
                        ws.cell(fila,cfg["verificacion"]).value=f"ERROR DE NAVEGACION: {problema}"
                        errores+=1
                        self.log(f"  -> {problema}")
                    wb.save(salida)
                    continue

                marca_web=combinar_marca_modelo(datos or {})
                color_web=(datos or {}).get("Color","")
                self.log(f"  Información de Objeto: Marca/Modelo={marca_web} | Color={color_web}")

                if not marca_web and not color_web:
                    ws.cell(fila,cfg["verificacion"]).value="DATOS DE OBJETO NO LEIDOS"
                    errores+=1
                    self.log("  -> DATOS DE OBJETO NO LEIDOS")
                else:
                    fallas=[]
                    if not coincide_marca_modelo(marca_excel, marca_web): fallas.append("MARCA/MODELO")
                    if not coincide_color(color_excel, color_web): fallas.append("COLOR")
                    if not fallas:
                        ws.cell(fila,cfg["verificacion"]).value="VERIFICADO SIN DOMINIO"
                        verificados+=1
                        self.log("  -> VERIFICADO SIN DOMINIO")
                    else:
                        ws.cell(fila,cfg["verificacion"]).value="NO COINCIDE: "+", ".join(fallas)
                        diferencias+=1
                        self.log("  -> DIFERENCIAS: "+", ".join(fallas))
                wb.save(salida)

            except Exception as e:
                errores+=1
                ws.cell(fila,cfg["verificacion"]).value=f"ERROR DE CONSULTA: {str(e)[:120]}"
                wb.save(salida)
                self.log(f"  -> ERROR: {e}")

        wb.save(salida)
        self.ultimo_archivo_salida = salida
        self.set_estado("Proceso finalizado.")
        resumen=(f"Finalizado. Hoja: {cfg['sheet_name']} | Procesadas: {total} | Verificadas: {verificados} | "
                 f"Con diferencias: {diferencias} | No verificables/no encontradas: {no_encontrados} | Errores: {errores}")
        self.log(resumen)
        self.log(f"Archivo de salida: {salida}")
        self.root.after(0,lambda: messagebox.showinfo("Verificación finalizada",resumen+f"\n\nSalida:\n{salida}"))


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
